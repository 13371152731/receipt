import copy
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT.parent))
from cert_demo.extraction import extract,evaluate,TEMPLATES

class ExtractionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.raw=json.loads((ROOT.parent/'ocr_benchmark/gpu_standard/ocr_results.json').read_text(encoding='utf-8'))
    def parsed(self,key): return extract(self.raw[key][0])
    def test_ocr_misses_remain_missing(self):
        r=evaluate({'parsed':self.parsed('T01')})
        self.assertIn('person_name',r['missing'])
        self.assertEqual(r['fields']['valid_to']['value'],'')
    def test_overlap_and_perspective_labels(self):
        for key,name,op,end in [('T06','田慧杰','高压电工作业','2027-01-31'),('T09','于汝良','低压电工作业','2029-02-16')]:
            r=evaluate({'parsed':self.parsed(key)})
            self.assertEqual(r['fields']['person_name']['value'],name)
            self.assertEqual(r['fields']['operation']['value'],op)
            self.assertEqual(r['fields']['valid_to']['value'],end)
            self.assertEqual(r['missing'],[])
    def test_initial_issue_date_not_valid_from(self):
        f=self.parsed('T06')['fields']
        self.assertEqual(f['initial_issue_date']['value'],'2018-03-07')
        self.assertEqual(f['valid_from']['value'],'2021-02-01')
        self.assertIn('initial_issue_date',TEMPLATES['electrical']['required'])
        self.assertNotIn('valid_from',TEMPLATES['electrical']['required'])
    def test_malformed_date_not_guessed(self):
        f=self.parsed('T03')['fields']
        self.assertEqual(f['valid_from']['value'],'')
        self.assertEqual(f['valid_to']['value'],'2025-07-18')
    def test_hcip_not_applicable_fields(self):
        r=evaluate({'parsed':self.parsed('T07')})
        self.assertEqual(r['status'],'ready')
        self.assertEqual(r['fields']['review_date']['state'],'na')
        self.assertEqual(r['fields']['valid_to']['value'],'2026-05-25')
        self.assertEqual(self.parsed('T08')['type'],'electrical')
    def test_month_precision_and_review_semantics(self):
        p=self.parsed('T02'); r=evaluate({'parsed':p})
        self.assertIn('review_date',r['needs_review'])
        r=evaluate({'parsed':p,'confirmed':['review_date']})
        self.assertNotIn('review_date',r['needs_review'])
    def test_confirmation_does_not_bypass_invalid_date(self):
        r=evaluate({'parsed':self.parsed('T08'),'corrections':{'valid_to':'2029-02-30'},'confirmed':['valid_to']})
        self.assertIn('valid_to',r['errors'])
        r=evaluate({'parsed':self.parsed('T08'),'corrections':{'initial_issue_date':'2030-01-01'},'confirmed':['initial_issue_date']})
        self.assertIn('initial_issue_date',r['errors'])
        r=evaluate({'parsed':self.parsed('T08'),'corrections':{'valid_to':'2020-01-01'},'confirmed':['valid_to']})
        self.assertIn('valid_to',r['errors'])
    def test_manual_source_preserves_original(self):
        p=self.parsed('T07'); before=copy.deepcopy(p)
        r=evaluate({'parsed':p,'corrections':{'person_name':'测试姓名'}})
        self.assertEqual(r['fields']['person_name']['raw_value'],'Huijie Tian')
        self.assertEqual(r['fields']['person_name']['source'],'manual')
        self.assertEqual(p,before)
    def test_expiry_and_review_risks_remain_independent_after_approval(self):
        from datetime import date
        r={'parsed':self.parsed('T09'),'corrections':{'valid_to':'2026-09-19'},'reviewed_at':'2026-09-14'}
        result=evaluate(r,today=date(2026,9,14))
        self.assertEqual(result['risk']['kind'],'expiring')
        self.assertEqual(result['review_risk']['kind'],'review_due')
        self.assertEqual(result['status'],'approved')
        self.assertEqual(len(result['risks']),2)
    def test_ambiguous_review_date_is_not_classified_automatically(self):
        from cert_demo.extraction import review_details
        self.assertEqual(review_details({'parsed':self.parsed('T02')})['next_due_date'],'')
        self.assertEqual(review_details({'parsed':self.parsed('T09')})['next_due_date'],'2026-02-16')
        self.assertEqual(review_details({'parsed':self.parsed('T09'),'corrections':{'review_date':'2028-01-01'}})['next_due_date'],'')

class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp=tempfile.TemporaryDirectory(prefix='cert-demo-test-')
        os.environ['CERT_DEMO_DATA']=cls.temp.name
        from cert_demo import app
        from fastapi.testclient import TestClient
        cls.a=app
        cls.client=TestClient(app.app,headers={'X-Demo-Request':'1'})
        cls.client.__enter__()
        cls.raw=json.loads((ROOT.parent/'ocr_benchmark/gpu_standard/ocr_results.json').read_text(encoding='utf-8'))
    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None,None,None)
        cls.temp.cleanup()
    def setUp(self):
        with self.a.connection() as c:
            c.execute('DELETE FROM records');c.execute('DELETE FROM audit')
            c.execute('UPDATE settings SET payload=? WHERE key=?',(json.dumps({'version':1,'templates':TEMPLATES}),'templates'))
        from PIL import Image
        fixture=Path(self.temp.name)/'fixture.png';Image.new('RGB',(100,100)).save(fixture)
        r={'id':'fixture','created_at':self.a.now(),'filename':'fixture.png','sample':None,'sha256':'test','image_path':str(fixture),'display_path':None,'job_status':'done','parsed':extract(self.raw['T07'][0]),'raw':self.raw['T07'][0],'corrections':{},'confirmed':[],'type_override':None,'ocr_s':.3,'version':1,'reviewed_at':None,'reviewer':'','note':''}
        self.a.write(r)
    def patch(self,**data):
        r=self.a.read('fixture')
        return self.client.patch('/api/records/fixture',json={'version':r['version'],**data})
    def test_approval_and_optimistic_lock(self):
        self.assertEqual(self.patch(approve=True).status_code,422)
        response=self.patch(approve=True,reviewer='测试审核员')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['status'],'approved')
        self.assertEqual(self.client.patch('/api/records/fixture',json={'version':1}).status_code,409)
    def test_delete_approved_record_hides_everywhere_and_preserves_audit(self):
        self.patch(approve=True,reviewer='测试')
        version=self.a.read('fixture')['version']
        self.assertEqual(self.client.request('DELETE','/api/records/fixture',json={'version':1}).status_code,409)
        response=self.client.request('DELETE','/api/records/fixture',json={'version':version})
        self.assertEqual(response.status_code,200)
        self.assertEqual(self.client.get('/api/state').json()['records'],[])
        self.assertEqual(self.client.get('/api/records/fixture').status_code,404)
        self.assertEqual(self.client.get('/api/records/fixture/image').status_code,404)
        self.assertEqual(self.client.post('/api/records/fixture/retry').status_code,404)
        export=self.client.post('/api/export',json={'format':'xlsx','ids':['fixture']})
        from openpyxl import load_workbook
        self.assertEqual(load_workbook(io.BytesIO(export.content)).active.max_row,1)
        with self.a.connection() as c:
            saved=json.loads(c.execute('SELECT payload FROM records').fetchone()[0])
            self.assertTrue(saved['deleted_at'])
            self.assertTrue(saved['reviewed_at'])
            self.assertTrue(Path(saved['image_path']).exists())
            self.assertEqual(c.execute("SELECT COUNT(*) FROM audit WHERE action='record_deleted'").fetchone()[0],1)
    def test_delete_running_blocked_and_origin_checked(self):
        r=self.a.read('fixture');r['job_status']='running';self.a.write(r)
        self.assertEqual(self.client.request('DELETE','/api/records/fixture',json={'version':1}).status_code,409)
        self.assertEqual(self.client.request('DELETE','/api/records/fixture',json={'version':1},headers={'Origin':'https://example.com'}).status_code,403)
        self.assertEqual(len(self.a.all_records()),1)
    def test_batch_delete_atomic_conflict_then_success(self):
        second=copy.deepcopy(self.a.read('fixture'));second['id']='second';self.a.write(second)
        items=[{'id':'fixture','version':1},{'id':'second','version':99}]
        self.assertEqual(self.client.post('/api/records/batch-delete',json={'records':items}).status_code,409)
        self.assertEqual(len(self.a.all_records()),2)
        items[1]['version']=1
        response=self.client.post('/api/records/batch-delete',json={'records':items})
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['deleted'],2)
        self.assertEqual(len(self.a.all_records()),0)
        with self.a.connection() as c:
            self.assertEqual(c.execute("SELECT COUNT(*) FROM audit WHERE action='record_deleted'").fetchone()[0],2)
    def test_engine_controls_busy_and_idempotent(self):
        old=copy.deepcopy(self.a.ENGINE)
        try:
            with patch.object(self.a,'OCR',None),patch.object(self.a,'OCR_PENDING',1):
                self.assertEqual(self.client.post('/api/engine/stop').status_code,409)
            with patch.object(self.a,'OCR',None),patch.object(self.a,'OCR_PENDING',0),patch.object(self.a.POOL,'submit') as submit:
                self.a.ENGINE['status']='idle'
                self.assertEqual(self.client.post('/api/engine/start').json()['status'],'loading')
                self.client.post('/api/engine/start')
                self.assertEqual(submit.call_count,1)
                self.assertEqual(self.client.post('/api/engine/stop').status_code,409)
                self.a.ENGINE['status']='idle'
                self.assertEqual(self.client.post('/api/engine/stop').json()['status'],'idle')
                self.a.ENGINE['status']='stopping'
                self.assertEqual(self.client.post('/api/engine/start').status_code,409)
        finally:self.a.ENGINE.clear();self.a.ENGINE.update(old)
    def test_batch_delete_invalid_running_and_unknown(self):
        for items in [[],[{'id':'fixture','version':1}]*2,[{'id':'fixture','version':True}]]:
            self.assertEqual(self.client.post('/api/records/batch-delete',json={'records':items}).status_code,422)
        self.assertEqual(self.client.post('/api/records/batch-delete',json={'records':[{'id':'fixture','version':1},{'id':'absent','version':1}]}).status_code,404)
        self.assertEqual(len(self.a.all_records()),1)
        r=self.a.read('fixture');r['job_status']='running';self.a.write(r)
        self.assertEqual(self.client.post('/api/records/batch-delete',json={'records':[{'id':'fixture','version':1}]}).status_code,409)
    def test_structured_review_validation_persistence_and_export(self):
        info={'last_review_date':'2025-01-01','next_due_date':'2028-01-01','proof_status':'checked','proof_note':'测试证明编号 001'}
        self.assertEqual(self.patch(review_details=info).status_code,422)
        r=self.a.read('fixture');r['parsed']=extract(self.raw['T09'][0]);self.a.write(r)
        response=self.patch(review_details=info,reviewer='测试核对员')
        self.assertEqual(response.status_code,200)
        self.assertEqual(response.json()['review_info'],info)
        self.assertEqual(self.a.read('fixture')['parsed']['fields']['review_date']['value'],'2026-02-16')
        for invalid in ['2028-02-30','2028-1-001']:
            self.assertEqual(self.patch(review_details={**info,'next_due_date':invalid},reviewer='测试').status_code,422)
        self.assertEqual(self.patch(review_details={**info,'next_due_date':'2024-01-01'},reviewer='测试').status_code,422)
        from openpyxl import load_workbook
        ws=load_workbook(io.BytesIO(self.client.post('/api/export',json={'format':'xlsx'}).content)).active
        row=dict(zip([c.value for c in ws[1]],[c.value for c in ws[2]]))
        self.assertEqual(row['下次应复审日期'],'2028-01-01')
        self.assertEqual(row['复审证明依据'],'测试证明编号 001')
    def test_missing_field_blocks_approval_and_raw_is_immutable(self):
        self.assertEqual(self.patch(changes={'person_name':''}).status_code,200)
        self.assertEqual(self.patch(approve=True,reviewer='测试').status_code,422)
        self.assertEqual(self.a.read('fixture')['parsed']['fields']['person_name']['value'],'Huijie Tian')
    def test_rule_change_reopens_review(self):
        self.patch(approve=True,reviewer='测试')
        required={k:list(v['required']) for k,v in TEMPLATES.items()}
        required['hcip'].remove('direction')
        response=self.client.put('/api/settings',json={'version':1,'required':required})
        self.assertEqual(response.status_code,200)
        self.assertIsNone(self.a.read('fixture')['reviewed_at'])
        self.assertEqual(self.client.put('/api/settings',json={'version':1,'required':required}).status_code,409)
    def test_export_missing_na_and_formula_safety(self):
        self.patch(changes={'person_name':'','direction':'=1+1'})
        response=self.client.post('/api/export',json={'format':'xlsx'})
        from openpyxl import load_workbook
        ws=load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(ws['E2'].value,'【未识别】')
        self.assertEqual(ws['E2'].fill.fgColor.rgb,'00FDE8E7')
        headers=[c.value for c in ws[1]]
        self.assertEqual(ws.cell(2,headers.index('复审信息')+1).value,'不适用')
        self.assertEqual(ws['J2'].value,"'=1+1")
        self.assertNotEqual(ws['J2'].data_type,'f')
        csv=self.client.post('/api/export',json={'format':'csv'}).content
        self.assertTrue(csv.startswith(b'\xef\xbb\xbf'))
    def test_upload_validation_and_deduplication(self):
        self.assertEqual(self.client.post('/api/upload',files={'file':('bad.png',b'not an image')}).status_code,415)
        from PIL import Image
        blob=io.BytesIO();Image.new('RGB',(10,10)).save(blob,format='PNG')
        with patch.object(self.a.POOL,'submit') as submit:
            first=self.client.post('/api/upload',files={'file':('../../new.png',blob.getvalue())})
            second=self.client.post('/api/upload',files={'file':('duplicate.png',blob.getvalue())})
            self.assertEqual(first.json()['record']['filename'],'new.png')
            self.assertFalse(first.json()['duplicate']);self.assertTrue(second.json()['duplicate'])
            self.assertEqual(submit.call_count,1)
    def test_cross_origin_mutation_rejected(self):
        self.assertEqual(self.client.post('/api/samples',headers={'Origin':'https://example.com'}).status_code,403)
        self.assertEqual(self.client.get('/api/records/nonexistent/image').status_code,404)

    def test_region_acceptance_preserves_raw_and_requires_confidence_review(self):
        run={'id':'run1','record_id':'fixture','version':1,'box':[1,2,80,30],
             'candidates':[{'id':'candidate1','text':'Test Name','confidence':.5,'values':{'person_name':'Test Name'}}]}
        with self.a.connection() as c:
            c.execute('INSERT OR REPLACE INTO region_runs VALUES (?,?,?)',('run1','fixture',json.dumps(run)))
        choice={'person_name':{'run_id':'run1','candidate_id':'candidate1'}}
        response=self.patch(region_selections=choice)
        self.assertEqual(response.status_code,200)
        field=response.json()['fields']['person_name']
        self.assertEqual(field['source'],'region_ocr')
        self.assertEqual(field['raw_value'],'Huijie Tian')
        self.assertEqual(field['state'],'review')
        self.assertEqual(self.patch(region_selections=choice).status_code,409)
        self.assertEqual(self.patch(approve=True,reviewer='测试').status_code,422)
        self.assertEqual(self.patch(confirmed=['person_name'],approve=True,reviewer='测试').status_code,200)

    def test_region_api_rejects_invalid_region_and_wrong_field(self):
        with patch.object(self.a.POOL,'submit') as submit:
            for target,box in [('person_name',[0,0,120,50]),('review_date',[0,0,80,50])]:
                response=self.client.post('/api/records/fixture/region',json={'version':1,'target':target,'box':box,'mode':'line'})
                self.assertEqual(response.status_code,422)
            self.assertEqual(submit.call_count,0)

    def test_date_policy_migration_preserves_user_edits(self):
        r=self.a.read('fixture');r['parsed']=extract(self.raw['T06'][0]);r['raw']=self.raw['T06'][0]
        r['parsed']['fields'].pop('initial_issue_date');r['corrections']={'valid_from':'2021-02-02'};r['reviewed_at']='old';self.a.write(r)
        legacy=copy.deepcopy(TEMPLATES)
        legacy['electrical']['applicable'].remove('initial_issue_date')
        legacy['electrical']['required']=['valid_from' if k=='initial_issue_date' else k for k in legacy['electrical']['required']]
        with self.a.connection() as c:c.execute("UPDATE settings SET payload=? WHERE key='templates'",(json.dumps({'version':4,'templates':legacy}),))
        self.a.migrate_date_policy()
        r=self.a.read('fixture');checked=self.a.present(r)
        self.assertEqual(checked['fields']['initial_issue_date']['value'],'2018-03-07')
        self.assertEqual(checked['fields']['valid_from']['value'],'2021-02-02')
        self.assertIsNone(r['reviewed_at'])
        self.assertEqual(self.a.settings()['version'],5)
        self.a.migrate_date_policy();self.assertEqual(self.a.settings()['version'],5)

class RegionTests(unittest.TestCase):
    def test_coordinates_reject_outside_nan_and_tiny(self):
        from cert_demo.region import validate_box
        for box in [[-1,0,90,30],[0,0,900,30],[0,0,float('nan'),30],[0,0,2,2]]:
            with self.assertRaises(ValueError):validate_box(box,100,100)
        self.assertEqual(validate_box([10,10,95,30],100,100),[10,10,95,30])
    def test_explicit_date_range_and_initial_date(self):
        from cert_demo.region import candidate_values
        self.assertEqual(candidate_values('valid_to','有效期限 2019.03.18至2025.03.18'),{'valid_from':'2019-03-18','valid_to':'2025-03-18'})
        self.assertEqual(candidate_values('initial_issue_date','初领日期 2019.03.18'),{'initial_issue_date':'2019-03-18'})
        self.assertEqual(candidate_values('valid_to','2022.03.18 2025.03.18'),{})
    def test_other_field_line_is_not_a_name_candidate(self):
        from cert_demo.region import candidate_values
        self.assertEqual(candidate_values('person_name','电工作业'),{})
        self.assertEqual(candidate_values('person_name','SAWS'),{})
        self.assertEqual(candidate_values('person_name','姓名：马传荣'),{'person_name':'马传荣'})

if __name__=='__main__': unittest.main()
