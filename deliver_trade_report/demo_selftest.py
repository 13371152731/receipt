"""trade_report 模块自测：直接运行 `python demo_selftest.py` 即可验证。"""

from decimal import Decimal
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from trade_report import (  # noqa: E402
    MONTHLY_TRADE_HEADERS, OUR_NAMES, OUR_TAX_ID, build_trade_report,
    export_trade_workbook, trade_direction,
)


def sample(kind, ledger, net, tax, total, day, side, rid='', buyer=False):
    party = {'buyer_name': '济南翱众信息系统有限公司', 'buyer_tax_id': OUR_TAX_ID} if buyer else {
        'seller_name': '济南翱众信息系统有限公司', 'seller_tax_id': OUR_TAX_ID}
    return dict(invoice_date=day, ledger=ledger, kind=kind, net=net, tax=tax, total=total, id=rid, **party)


SAMPLES = [
    # 公司进项（我方是购买方）
    sample('special', 'company', '1000.00', '130.00', '1130.00', '2026-03-05', 'seller', 'in-1', buyer=True),
    sample('special', 'company', '2000.00', '260.00', '2260.00', '2026-03-11', 'seller', 'in-2', buyer=True),
    # 公司销项（我方是销售方）
    sample('special', 'company', '5000.00', '650.00', '5650.00', '2026-03-20', 'buyer', 'out-1'),
    # 销项红冲（负数）
    sample('ordinary', 'company', '-100.00', '-13.00', '-113.00', '2026-03-21', 'buyer', 'out-2'),
    # 进销项待确认：双方都不含我方
    dict(invoice_date='2026-03-22', ledger='company', kind='special', net='10.00', tax='1.30',
         total='11.30', buyer_name='甲公司', seller_name='乙公司', id='unk-1'),
    # 报销普票（计入「报销普票总额」）
    sample('ordinary', 'expense', '300.00', '0.00', '300.00', '2026-03-08', 'buyer', 'exp-1', buyer=True),
    # 报销专票（计入「报销专票份数」）
    sample('special', 'expense', '400.00', '52.00', '452.00', '2026-03-09', 'buyer', 'exp-2', buyer=True),
    # 2 月：只有报销，仍应出现一行进销项月报
    sample('ordinary', 'expense', '60.00', '0.00', '60.00', '2026-02-15', 'buyer', 'exp-3', buyer=True),
]


def main():
    assert trade_direction({'buyer_tax_id': OUR_TAX_ID}) == 'input'
    assert trade_direction({'seller_name': '济南翱众信息系统有限公司'}) == 'output'
    assert trade_direction({'buyer_name': '甲公司', 'seller_name': '乙公司'}) == 'unknown'

    rows = build_trade_report(SAMPLES, 2026)
    assert [r['month'] for r in rows] == ['2026-03', '2026-02'], rows

    mar = rows[0]
    assert mar['input_net'] == '3000.00' and mar['input_tax'] == '390.00' and mar['input_total'] == '3390.00', mar
    assert mar['output_net'] == '4900.00' and mar['output_tax'] == '637.00' and mar['output_total'] == '5537.00', mar
    assert mar['tax_diff'] == '-247.00', mar
    assert mar['expense_total'] == '300.00', mar
    assert mar['unknown_count'] == 1, mar
    assert mar['expense_special_count'] == 1, mar
    assert '2份进项发票' in mar['composition'] and '1份销项负数发票' in mar['composition'], mar

    feb = rows[1]
    assert feb['expense_total'] == '60.00' and feb['input_count'] == 0, feb

    assert build_trade_report(SAMPLES, 2025) == []

    payload = export_trade_workbook(rows)
    from io import BytesIO
    from openpyxl import load_workbook
    wb = load_workbook(BytesIO(payload))
    assert wb.sheetnames == ['进销项月报'], wb.sheetnames
    ws = wb['进销项月报']
    assert [c.value for c in ws[1]] == MONTHLY_TRADE_HEADERS
    assert ws.cell(2, 1).value == '2026-03' and ws.cell(2, 4).value == 3390.0

    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'demo_output.xlsx')
    with open(out, 'wb') as fh:
        fh.write(payload)
    print('全部断言通过；示例文件已生成：', out)


if __name__ == '__main__':
    main()