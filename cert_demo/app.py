"""Loopback-only certificate review demo with real, local PaddleOCR."""
import argparse
import asyncio
import contextlib
from concurrent.futures import ThreadPoolExecutor
from contextlib import asynccontextmanager, contextmanager
import copy
import csv
import gc
from datetime import datetime
import hashlib
import io
import json
import logging
import os
from pathlib import Path
import sqlite3
import sys
import threading
import time
import uuid

from fastapi import FastAPI, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, JSONResponse, Response
from fastapi.staticfiles import StaticFiles
from starlette.middleware.trustedhost import TrustedHostMiddleware
from PIL import Image, ImageOps, UnidentifiedImageError

ROOT=Path(__file__).resolve().parent
sys.path.insert(0,str(ROOT.parent))
from cert_demo.extraction import FIELD_DEFS, LABELS, TEMPLATES, evaluate, extract, review_details
from cert_demo.region import validate_box, candidate_values
from cert_demo import operations
from invoice_demo import service as invoice_service
DATA=Path(os.environ.get('CERT_DEMO_DATA',str(ROOT/'data')))
DATA.mkdir(parents=True,exist_ok=True)
(DATA/'images').mkdir(exist_ok=True)
DB=DATA/'demo.sqlite3'
LOCK=threading.RLock()
POOL=ThreadPoolExecutor(max_workers=1,thread_name_prefix='local-ocr')
ENGINE={'status':'idle','message':'首次识别时加载本地模型','device':None,'initialization_s':None}
OCR=None
OCR_PENDING=0

def submit_ocr(fn,*args):
    global OCR_PENDING
    with LOCK:
        if ENGINE['status']=='stopping':raise HTTPException(409,'模型正在关闭，请稍后重试')
        OCR_PENDING+=1
        def run():
            global OCR_PENDING
            try:return fn(*args)
            finally:
                with LOCK:OCR_PENDING-=1
        try:return POOL.submit(run)
        except Exception:
            OCR_PENDING-=1
            raise

def engine_state():
    with LOCK:return {**ENGINE,'pending_tasks':OCR_PENDING,'loaded':OCR is not None}

def start_engine():
    try:engine()
    except Exception as e:
        log.exception('Model startup failed')
        ENGINE.update(status='error',message=str(e)[:250],device=None)

def stop_engine():
    global OCR
    try:
        OCR=None
        gc.collect()
        paddle=sys.modules.get('paddle')
        if paddle is not None and paddle.device.get_device().startswith('gpu'):paddle.device.cuda.empty_cache()
        ENGINE.update(status='idle',message='模型已关闭；上传或框选识别时会自动启动',device=None,initialization_s=None)
    except Exception:
        log.exception('Model cleanup failed')
        ENGINE.update(status='error',message='模型已卸载，但运行时缓存清理失败',device=None)
MAX_UPLOAD=15*1024*1024
MAX_PIXELS=20_000_000
Image.MAX_IMAGE_PIXELS=MAX_PIXELS
log=logging.getLogger('certificate-demo')

def now(): return datetime.now().astimezone().isoformat(timespec='seconds')
@contextmanager
def connection():
    c=sqlite3.connect(DB,timeout=20); c.row_factory=sqlite3.Row
    try:
        with c: yield c
    finally: c.close()
def initialize():
    with connection() as c:
        c.execute('CREATE TABLE IF NOT EXISTS records (id TEXT PRIMARY KEY, created TEXT NOT NULL, payload TEXT NOT NULL)')
        c.execute('CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, payload TEXT NOT NULL)')
        c.execute('CREATE TABLE IF NOT EXISTS audit (id INTEGER PRIMARY KEY, record_id TEXT, at TEXT, action TEXT, payload TEXT)')
        c.execute('CREATE TABLE IF NOT EXISTS region_runs (id TEXT PRIMARY KEY, record_id TEXT, payload TEXT)')
        c.execute('INSERT OR IGNORE INTO settings VALUES (?,?)',('templates',json.dumps({'version':1,'date_policy':2,'templates':TEMPLATES},ensure_ascii=False)))
    migrate_date_policy()
    operations.initialize(sys.modules[__name__])
    # A restart never silently leaves jobs looking as if they were running.
    for r in all_records():
        if r['job_status'] in ['queued','running']:
            r['job_status']='error'; r['error']='服务重启中断了识别，请点击重新识别'; r['version']+=1; write(r)

def migrate_date_policy():
    old=settings()
    if old.get('date_policy')==2:return
    backup=DATA/'before-date-policy-v2.sqlite3'
    if not backup.exists():
        with connection() as source:
            dest=sqlite3.connect(backup)
            try:source.backup(dest)
            finally:dest.close()
    new=copy.deepcopy(old);new['version']+=1;new['date_policy']=2
    for typ in ['electrical','unknown']:
        new['templates'][typ]['applicable']=TEMPLATES[typ]['applicable']
    keys=new['templates']['electrical']['required']
    new['templates']['electrical']['required']=list(dict.fromkeys('initial_issue_date' if k=='valid_from' else k for k in keys))
    with connection() as c:
        c.execute('UPDATE settings SET payload=? WHERE key=?',(json.dumps(new,ensure_ascii=False),'templates'))
        for row in c.execute('SELECT payload FROM records').fetchall():
            r=json.loads(row['payload'])
            if r.get('raw') and r.get('parsed'):
                r['parsed']['fields']['initial_issue_date']=extract(r['raw'])['fields']['initial_issue_date']
            if (r.get('type_override') or (r.get('parsed') or {}).get('type'))!='hcip':
                r['reviewed_at']=None;r['version']+=1
            c.execute('UPDATE records SET payload=? WHERE id=?',(json.dumps(r,ensure_ascii=False),r['id']))
        c.execute('INSERT INTO audit(record_id,at,action,payload) VALUES (?,?,?,?)',(None,now(),'date_policy_updated',json.dumps({'business_start':'initial_issue_date','original_valid_from_preserved':True})))
def all_records():
    with connection() as c: return [record for r in c.execute('SELECT payload FROM records ORDER BY created DESC,id') if not (record:=json.loads(r['payload'])).get('deleted_at')]
def read(rid):
    with connection() as c: r=c.execute('SELECT payload FROM records WHERE id=?',(rid,)).fetchone()
    if not r: raise HTTPException(404,'记录不存在')
    record=json.loads(r['payload'])
    if record.get('deleted_at'): raise HTTPException(404,'记录已删除')
    return record
def write(r):
    with connection() as c: c.execute('INSERT OR REPLACE INTO records VALUES (?,?,?)',(r['id'],r['created_at'],json.dumps(r,ensure_ascii=False)))
def audit(rid,action,payload):
    with connection() as c: c.execute('INSERT INTO audit(record_id,at,action,payload) VALUES (?,?,?,?)',(rid,now(),action,json.dumps(payload,ensure_ascii=False)))
def settings():
    with connection() as c: return json.loads(c.execute('SELECT payload FROM settings WHERE key=?',('templates',)).fetchone()['payload'])
def present(r,details=False):
    v=evaluate(r,settings()['templates'],warning_days=max(operations.policy(sys.modules[__name__])['days']))
    result={k:val for k,val in r.items() if k not in ['raw','parsed','image_path','display_path']}
    result.update(v)
    result['image_url']=f'/api/records/{r["id"]}/image'
    result['original_url']=f'/api/records/{r["id"]}/image?original=true'
    if details:
        result['raw']=r.get('raw')
        with connection() as c: result['audit']=[dict(x) for x in c.execute('SELECT at,action,payload FROM audit WHERE record_id=? ORDER BY id DESC LIMIT 30',(r['id'],))]
    return result

def synchronize_ocr():
    """GPU timing needs synchronization; CPU inference is already synchronous."""
    paddle=sys.modules.get('paddle')
    if paddle is not None and paddle.device.get_device().startswith('gpu'):
        paddle.device.cuda.synchronize()

def engine():
    global OCR
    if OCR is not None: return OCR
    ENGINE.update(status='loading',message='正在加载本地 OCR 模型')
    t=time.perf_counter()
    os.environ.setdefault('PADDLE_PDX_CACHE_HOME',str(ROOT.parent/'models'/'paddlex-cache'))
    os.environ['PADDLE_PDX_DISABLE_MODEL_SOURCE_CHECK']='True'
    os.environ['USE_TORCH']='0'
    sys.modules['torch']=None  # Keep unrelated Torch CUDA DLLs out of Paddle process.
    if os.environ.get('CERT_OCR_DEVICE')=='cpu' and sys.platform=='linux':
        cpuinfo=Path('/proc/cpuinfo').read_text()
        if ' avx ' not in cpuinfo:
            raise RuntimeError('虚拟 CPU 未启用 AVX，请在 Proxmox 将处理器类型改为 host，关机后重新开机')
    from paddleocr import PaddleOCR
    import paddle
    config_path=ROOT.parent/'ocr_benchmark'/'gpu_standard'/'metadata.json'
    cfg=json.loads(config_path.read_text(encoding='utf-8'))['config']
    if os.environ.get('CERT_OCR_MODEL_ROOT'):
        for key in [k for k in cfg if k.endswith('_model_dir')]:
            cfg[key]=str(Path(os.environ['CERT_OCR_MODEL_ROOT'])/cfg[key].replace('\\','/').rsplit('/',1)[-1])
    cfg['device']=os.environ.get('CERT_OCR_DEVICE',cfg['device'])
    cfg['cpu_threads']=int(os.environ.get('CERT_OCR_CPU_THREADS',cfg.get('cpu_threads',4)))
    if cfg['device']=='cpu':cfg['enable_mkldnn']=os.environ.get('CERT_OCR_MKLDNN','1')=='1'
    for key,path in cfg.items():
        if key.endswith('_model_dir') and not Path(path).is_dir():
            raise RuntimeError('本地模型目录不存在，请按 README 配置 CERT_OCR_MODEL_ROOT')
    if cfg['device'].startswith('gpu') and (not paddle.is_compiled_with_cuda() or paddle.device.cuda.device_count()<1):
        raise RuntimeError('当前 Demo 配置需要可用的 Paddle CUDA 环境，请使用已验证的本地 Python')
    OCR=PaddleOCR(**cfg)
    synchronize_ocr()
    device=paddle.device.cuda.get_device_name(0) if cfg['device'].startswith('gpu') else 'CPU'
    ENGINE.update(status='ready',message='本地 OCR 已就绪',device=device,initialization_s=round(time.perf_counter()-t,3))
    return OCR

def native(x):
    if hasattr(x,'tolist'): return x.tolist()
    if isinstance(x,dict): return {str(k):native(v) for k,v in x.items()}
    if isinstance(x,(list,tuple)): return [native(v) for v in x]
    return x

def recognize_region(r,box,target,mode):
    ocr=engine()
    import cv2
    import numpy as np
    import paddle
    pixels=cv2.imdecode(np.frombuffer(Path(r.get('display_path') or r['image_path']).read_bytes(),np.uint8),cv2.IMREAD_COLOR)
    x1,y1,x2,y2=box
    crop=pixels[y1:y2,x1:x2].copy()
    synchronize_ocr();t=time.perf_counter()
    candidates=[]
    if mode=='line':
        # Reuse the installed PaddleX recognition model, bypassing detection.
        gray=cv2.cvtColor(crop,cv2.COLOR_BGR2GRAY)
        enhanced=cv2.cvtColor(cv2.createCLAHE(clipLimit=2.0,tileGridSize=(8,8)).apply(gray),cv2.COLOR_GRAY2BGR)
        outputs=list(ocr.paddlex_pipeline.text_rec_model([crop,enhanced]))
        texts=[(res['rec_text'],float(res['rec_score']),label) for res,label in zip(outputs,['原色区域','灰度增强区域'])]
    else:
        outputs=list(ocr.predict(crop,use_doc_orientation_classify=False,use_doc_unwarping=False))
        texts=[]
        for res in outputs:
            scores=native(res['rec_scores'])
            texts.extend((text,float(score),f'区域文字行 {i+1}') for i,(text,score) in enumerate(zip(res['rec_texts'],scores)))
        texts=texts[:30]
    synchronize_ocr();elapsed=time.perf_counter()-t
    for text,score,label in texts:
        text=str(text)
        candidates.append({'id':uuid.uuid4().hex,'text':text,'confidence':score,'variant':label,
                           'values':candidate_values(target,text)})
    result={'id':uuid.uuid4().hex,'record_id':r['id'],'version':r['version'],'target':target,'box':box,'mode':mode,
            'ocr_s':round(elapsed,4),'candidates':candidates,'at':now()}
    with connection() as c:c.execute('INSERT INTO region_runs VALUES (?,?,?)',(result['id'],r['id'],json.dumps(result,ensure_ascii=False)))
    audit(r['id'],'region_ocr',{'run_id':result['id'],'target':target,'box':box,'ocr_s':result['ocr_s']})
    return result

def apply_region_selections(r,selections):
    if not isinstance(selections,dict):raise HTTPException(422,'区域候选选择无效')
    allowed=settings()['templates'][r.get('type_override') or r['parsed']['type']]['applicable']
    for key,choice in selections.items():
        if key not in allowed or not isinstance(choice,dict):raise HTTPException(422,'区域字段不适用')
        with connection() as c:row=c.execute('SELECT payload FROM region_runs WHERE id=? AND record_id=?',(choice.get('run_id'),r['id'])).fetchone()
        if not row:raise HTTPException(422,'区域候选不存在或不属于该图片')
        run=json.loads(row['payload'])
        if run['version']!=r['version']:raise HTTPException(409,'区域候选已过时，请重新框选识别')
        candidate=next((x for x in run['candidates'] if x['id']==choice.get('candidate_id')),None)
        if not candidate or not candidate['values'].get(key):raise HTTPException(422,'该候选没有可填入的字段值')
        r.setdefault('region_overrides',{})[key]={'value':candidate['values'][key],'confidence':candidate['confidence'],
            'evidence':[{'text':candidate['text'],'box':run['box'],'index':-1}],
            'run_id':run['id'],'candidate_id':candidate['id']}
        r['corrections'].pop(key,None)

def process_record(rid):
    try:
        with LOCK:
            r=read(rid); r['job_status']='running'; r['version']+=1; r['error']=None; write(r)
        ocr=engine()
        import cv2
        import numpy as np
        import paddle
        t=time.perf_counter()
        pixels=cv2.imdecode(np.frombuffer(Path(r['image_path']).read_bytes(),np.uint8),cv2.IMREAD_COLOR)
        decode_s=time.perf_counter()-t
        synchronize_ocr(); t=time.perf_counter()
        res=list(ocr.predict(pixels))
        synchronize_ocr(); elapsed=time.perf_counter()-t
        if len(res)!=1: raise RuntimeError('本次图片没有返回单页识别结果')
        result=res[0]
        raw={key:native(result[key]) for key in ['rec_texts','rec_scores','rec_polys']}
        angle=int(result.get('doc_preprocessor_res',{}).get('angle',0))
        raw['orientation_angle']=angle
        # Store the OCR-aligned image so field boxes match the visible coordinates.
        processed=result.get('doc_preprocessor_res',{}).get('output_img')
        if processed is None: processed=pixels
        display=DATA/'images'/f'{rid}-aligned.png'
        success,encoded=cv2.imencode('.png',processed)
        if not success: raise RuntimeError('无法保存字段对照图')
        display.write_bytes(encoded.tobytes())
        parsed=extract(raw)
        with LOCK:
            r=read(rid); r.update(raw=raw,parsed=parsed,job_status='done',ocr_s=round(elapsed,4),decode_s=round(decode_s,4),display_path=str(display),processed_at=now(),reviewed_at=None,confirmed=[])
            r['version']+=1; write(r); audit(rid,'ocr_completed',{'ocr_s':r['ocr_s'],'model':'PP-OCRv5 server','angle':angle})
    except Exception as e:
        log.exception('OCR failed for %s',rid)
        if OCR is None: ENGINE.update(status='error',message=str(e)[:250])
        with LOCK:
            r=read(rid); r.update(job_status='error',error=str(e)[:500]); r['version']+=1; write(r)

def create_record(content,filename,sample=None):
    if len(content)>MAX_UPLOAD: raise HTTPException(413,'单张图片不能超过 15 MB')
    try:
        with Image.open(io.BytesIO(content)) as img:
            if img.format not in ['PNG','JPEG','WEBP']: raise HTTPException(415,'仅支持 PNG、JPG、WEBP 图片')
            if img.width*img.height>MAX_PIXELS: raise HTTPException(413,'图片分辨率过大，最多 2000 万像素')
            image=ImageOps.exif_transpose(img).convert('RGB')
            width,height=image.size
    except (UnidentifiedImageError,OSError,Image.DecompressionBombError):
        raise HTTPException(415,'文件不是可读取的图片')
    digest=hashlib.sha256(content).hexdigest()
    with LOCK:
        if ENGINE['status']=='stopping':raise HTTPException(409,'模型正在关闭，请稍后上传')
        existing=next((r for r in all_records() if r['sha256']==digest),None)
        if existing: return existing,True
        rid=uuid.uuid4().hex
        image_path=DATA/'images'/f'{rid}.png'
        image.save(image_path)
        r={'id':rid,'filename':Path(filename.replace('\\','/')).name[:150],'sample':sample,'sha256':digest,'image_path':str(image_path),'display_path':None,'width':width,'height':height,'created_at':now(),'job_status':'queued','error':None,'parsed':None,'raw':None,'corrections':{},'confirmed':[],'type_override':None,'ocr_s':None,'version':1,'reviewed_at':None,'reviewer':'','note':''}
        write(r); audit(rid,'uploaded',{'filename':r['filename'],'sample':sample})
        submit_ocr(process_record,rid)
    return r,False

@asynccontextmanager
async def lifespan(app):
    initialize()
    app.state.invoices.initialize()
    async def scan():
        while True:
            try:
                await asyncio.to_thread(operations.sync_alerts, sys.modules[__name__])
            except Exception:
                log.exception('Alert scan failed')
            await asyncio.sleep(60)
    scan_task=asyncio.create_task(scan())
    yield
    scan_task.cancel()
    with contextlib.suppress(asyncio.CancelledError):
        await scan_task
    POOL.shutdown(wait=False,cancel_futures=True)

app=FastAPI(title='证书提取与审核 Demo',lifespan=lifespan,docs_url=None,redoc_url=None)
app.add_middleware(TrustedHostMiddleware,allowed_hosts=os.environ.get('APP_ALLOWED_HOSTS','127.0.0.1,localhost,testserver').split(','))

@app.middleware('http')
async def local_only(request:Request,call_next):
    if request.method not in ['GET','HEAD','OPTIONS']:
        origin=request.headers.get('origin')
        expected=f'{request.url.scheme}://{request.headers.get("host", "")}'
        if (origin and origin!=expected) or request.headers.get('x-demo-request')!='1':
            return JSONResponse({'detail':'请求来源不匹配，请在本地 Demo 页面操作'},status_code=403)
    response=await call_next(request)
    response.headers['Cache-Control']='no-store'
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['Referrer-Policy']='no-referrer'
    response.headers['Content-Security-Policy']="default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' blob: data:; connect-src 'self'; frame-ancestors 'none'"
    return response

@app.get('/api/state')
def state():
    with LOCK:
        records=all_records()
        historical=operations.historical_ids(sys.modules[__name__],records)
        people={p['id']:p for p in operations.load_rows(sys.modules[__name__],'people')}
        rows=[]
        for r in records:
            person=people.get(r.get('person_id'))
            rows.append({**present(r),'historical':r['id'] in historical,'person':person})
        return {'records':rows,'engine':engine_state(),'settings':settings(),'fields':[{'key':k,'label':v} for k,v in FIELD_DEFS],'today':datetime.now().date().isoformat()}

@app.post('/api/engine/start')
def start_model():
    with LOCK:
        if ENGINE['status']=='stopping':raise HTTPException(409,'模型正在关闭，请稍后启动')
        if ENGINE['status'] not in ['loading','ready']:
            ENGINE.update(status='loading',message='正在加载本地 OCR 模型')
            POOL.submit(start_engine)
    return engine_state()

@app.post('/api/engine/stop')
def stop_model():
    with LOCK:
        if OCR_PENDING or ENGINE['status']=='loading' or any(r['job_status'] in ['queued','running'] for r in all_records()):
            raise HTTPException(409,'模型加载或识别任务尚未完成，请稍后关闭')
        if ENGINE['status']=='stopping':return engine_state()
        if OCR is None:
            ENGINE.update(status='idle',message='模型未启动',device=None)
        else:
            ENGINE.update(status='stopping',message='正在卸载模型并清理可释放的显存缓存')
            POOL.submit(stop_engine)
    return engine_state()

@app.get('/api/records/{rid}')
def detail(rid:str): return present(read(rid),True)

@app.get('/api/records/{rid}/image')
def image(rid:str,original:bool=False):
    r=read(rid); path=r['image_path'] if original else r.get('display_path') or r['image_path']
    return FileResponse(path,media_type='image/png')

@app.post('/api/upload')
async def upload(file:UploadFile=File(...)):
    content=await file.read(MAX_UPLOAD+1); await file.close()
    r,duplicate=create_record(content,file.filename or '证书.png')
    return {'record':present(r),'duplicate':duplicate}

@app.post('/api/samples')
def samples():
    meta=json.loads((ROOT.parent/'ocr_benchmark'/'gpu_standard'/'metadata.json').read_text(encoding='utf-8'))
    added=duplicates=0
    for s in meta['samples']:
        path=Path(s['source_path'])
        if not path.is_file(): raise HTTPException(404,f'样例图片不存在：{s["filename"]}')
        _,duplicate=create_record(path.read_bytes(),s['filename'],s['id'])
        if duplicate: duplicates+=1
        else: added+=1
    return {'added':added,'duplicates':duplicates}

@app.post('/api/records/batch-delete')
async def batch_delete(request:Request):
    body=await request.json()
    items=body.get('records') if isinstance(body,dict) else None
    if not isinstance(items,list) or not 1<=len(items)<=500 or any(not isinstance(x,dict) or not isinstance(x.get('id'),str) or type(x.get('version')) is not int for x in items):
        raise HTTPException(422,'请选择 1 至 500 条有效记录')
    if len({x['id'] for x in items})!=len(items):raise HTTPException(422,'不能重复选择同一记录')
    with LOCK:
        records=[]
        for item in items:
            r=read(item['id'])
            if r['job_status'] in ['queued','running']:raise HTTPException(409,'所选记录正在识别，本次未删除任何条目')
            if r['version']!=item['version']:raise HTTPException(409,'所选记录已更新，请重新勾选；本次未删除任何条目')
            records.append(r)
        at=now();batch_id=uuid.uuid4().hex
        with connection() as c:
            for r in records:
                r['deleted_at']=at;r['version']+=1
                c.execute('UPDATE records SET payload=? WHERE id=?',(json.dumps(r,ensure_ascii=False),r['id']))
                c.execute('INSERT INTO audit(record_id,at,action,payload) VALUES (?,?,?,?)',(r['id'],at,'record_deleted',json.dumps({'batch_id':batch_id,'version':r['version']})))
    return {'deleted':len(records),'ids':[r['id'] for r in records]}

@app.delete('/api/records/{rid}')
async def delete_record(rid:str,request:Request):
    body=await request.json()
    if not isinstance(body,dict): raise HTTPException(422,'删除请求不合法')
    with LOCK:
        r=read(rid)
        if r['job_status'] in ['queued','running']: raise HTTPException(409,'证书正在识别，请完成后再删除')
        if body.get('version')!=r['version']: raise HTTPException(409,'记录已更新，请刷新后重新确认删除')
        r['deleted_at']=now();r['version']+=1
        with connection() as c:
            c.execute('UPDATE records SET payload=? WHERE id=?',(json.dumps(r,ensure_ascii=False),rid))
            c.execute('INSERT INTO audit(record_id,at,action,payload) VALUES (?,?,?,?)',(rid,now(),'record_deleted',json.dumps({'version':r['version']},ensure_ascii=False)))
    return {'deleted':True,'id':rid}

@app.post('/api/records/{rid}/retry')
def retry(rid:str):
    with LOCK:
        if ENGINE['status']=='stopping':raise HTTPException(409,'模型正在关闭，请稍后重试')
        r=read(rid)
        if r['job_status'] in ['queued','running']: raise HTTPException(409,'该图片正在识别')
        r['job_status']='queued'; r['version']+=1; r['reviewed_at']=None; write(r)
        audit(rid,'ocr_retry',{'manual_corrections_preserved':True}); submit_ocr(process_record,rid)
    return present(r)

@app.post('/api/records/{rid}/region')
async def region_ocr(rid:str,request:Request):
    body=await request.json()
    with LOCK:
        r=read(rid)
        if r['job_status']!='done' or body.get('version')!=r['version']:raise HTTPException(409,'记录已更新或正在识别，请重新打开')
        target=body.get('target');mode=body.get('mode','line')
        if target not in present(r)['fields'] or not present(r)['fields'][target]['applicable']:raise HTTPException(422,'该字段不支持框选')
        if mode not in ['line','block']:raise HTTPException(422,'识别模式无效')
        with Image.open(r.get('display_path') or r['image_path']) as im:
            try:box=validate_box(body.get('box'),*im.size)
            except ValueError as e:raise HTTPException(422,str(e))
    try:return await asyncio.wrap_future(submit_ocr(recognize_region,r,box,target,mode))
    except HTTPException:raise
    except Exception as e:
        log.exception('Region OCR failed')
        raise HTTPException(500,'局部识别失败：'+str(e)[:200])

@app.patch('/api/records/{rid}')
async def update(rid:str,request:Request):
    data=await request.json()
    with LOCK:
        r=read(rid)
        if r['job_status']!='done': raise HTTPException(409,'识别完成后才能修改字段')
        if data.get('version')!=r['version']: raise HTTPException(409,'记录已更新，请重新打开后修改')
        typ=data.get('type',r.get('type_override') or r['parsed']['type'])
        if typ not in TEMPLATES: raise HTTPException(422,'不支持的证书类型')
        changes=data.get('changes',{}); confirms=data.get('confirmed',[])
        if not isinstance(changes,dict) or any(k not in LABELS or not isinstance(v,str) or len(v)>200 for k,v in changes.items()):
            raise HTTPException(422,'字段值不合法')
        if not isinstance(confirms,list) or any(k not in LABELS for k in confirms): raise HTTPException(422,'字段确认列表不合法')
        before={'corrections':copy.deepcopy(r['corrections']),'confirmed':r['confirmed'],'type_override':r['type_override'],'reviewed_at':r['reviewed_at'],'review_details':copy.deepcopy(r.get('review_details',{}))}
        if 'review_details' in data:
            info=data['review_details']
            expected={'last_review_date','next_due_date','proof_status','proof_note'}
            if not isinstance(info,dict) or set(info)!=expected or any(not isinstance(v,str) for v in info.values()):raise HTTPException(422,'复审信息结构不合法')
            info={k:v.strip() for k,v in info.items()}
            if info['proof_status'] not in ['unknown','pending','checked'] or len(info['proof_note'])>500:raise HTTPException(422,'复审证明状态或说明不合法')
            for key in ['last_review_date','next_due_date']:
                if info[key]:
                    try:
                        if len(info[key])!=10:raise ValueError()
                        if datetime.strptime(info[key],'%Y-%m-%d').date().isoformat()!=info[key]:raise ValueError()
                    except ValueError:raise HTTPException(422,'复审日期须为有效的 YYYY-MM-DD 日期')
            if info['last_review_date'] and info['last_review_date']>datetime.now().date().isoformat():raise HTTPException(422,'上次复审日期不能晚于今天')
            if info['last_review_date'] and info['next_due_date'] and info['last_review_date']>=info['next_due_date']:raise HTTPException(422,'下次应复审日期须晚于上次复审日期')
            if info['proof_status']=='checked' and (not info['proof_note'] or not str(data.get('reviewer',r.get('reviewer',''))).strip()):raise HTTPException(422,'核对复审证明后请填写核对人（审核人）和证明依据')
            r['review_details']=info
        r['type_override']=typ
        apply_region_selections(r,data.get('region_selections',{}))
        r['corrections'].update({k:v.strip() for k,v in changes.items()})
        r['confirmed']=confirms; r['reviewed_at']=None
        r['reviewer']=str(data.get('reviewer',r.get('reviewer','')))[:80]
        r['note']=str(data.get('note',r.get('note','')))[:1000]
        if data.get('approve'):
            checked=evaluate(r,settings()['templates'])
            if checked['missing'] or checked['needs_review'] or checked['notes']:
                raise HTTPException(422,'仍有缺失或待核对字段，请补录或勾选已核对后再确认审核')
            if not r['reviewer'].strip(): raise HTTPException(422,'确认审核前请填写审核人')
            r['reviewed_at']=now()
        r['version']+=1; write(r)
        audit(rid,'approved' if data.get('approve') else 'fields_updated',{'before':before,'changes':changes,'review_details':r.get('review_details',{}),'region_selections':data.get('region_selections',{}),'confirmed':confirms,'type':typ,'reviewer':r['reviewer'],'note':r['note']})
    return present(r,True)

@app.put('/api/settings')
async def change_settings(request:Request):
    body=await request.json()
    with LOCK:
        old=settings()
        if body.get('version')!=old['version']: raise HTTPException(409,'字段规则已更新，请重新打开')
        required=body.get('required',{})
        if set(required)!=set(TEMPLATES): raise HTTPException(422,'字段模板不完整')
        new=copy.deepcopy(old); new['version']+=1
        for typ,keys in required.items():
            if not isinstance(keys,list) or len(keys)!=len(set(keys)) or any(k not in TEMPLATES[typ]['applicable'] for k in keys):
                raise HTTPException(422,'必填字段配置不合法')
            new['templates'][typ]['required']=keys
        # Rule changes invalidate earlier review decisions for affected types.
        affected={typ for typ in TEMPLATES if old['templates'][typ]['required']!=new['templates'][typ]['required']}
        with connection() as c:
            c.execute('UPDATE settings SET payload=? WHERE key=?',(json.dumps(new,ensure_ascii=False),'templates'))
            for row in c.execute('SELECT id,payload FROM records').fetchall():
                r=json.loads(row['payload']); typ=r.get('type_override') or (r.get('parsed') or {}).get('type','unknown')
                if typ in affected:
                    r['reviewed_at']=None; r['version']+=1
                    c.execute('UPDATE records SET payload=? WHERE id=?',(json.dumps(r,ensure_ascii=False),r['id']))
            c.execute('INSERT INTO audit(record_id,at,action,payload) VALUES (?,?,?,?)',(None,now(),'templates_updated',json.dumps({'before':old,'after':new},ensure_ascii=False)))
    return new

def spreadsheet_safe(value):
    text=str(value if value is not None else '')
    return "'"+text if text.lstrip().startswith(('=','+','-','@','\t','\r')) else text

@app.post('/api/export')
async def export(request:Request):
    body=await request.json(); selected=body.get('ids')
    if selected is not None and (not isinstance(selected,list) or any(not isinstance(i,str) for i in selected)):
        raise HTTPException(422,'导出范围不合法')
    with LOCK:
        source_records=all_records()
        historical=operations.historical_ids(sys.modules[__name__],source_records)
        people={p['id']:p for p in operations.load_rows(sys.modules[__name__],'people')}
        rows=[present(r) for r in source_records if selected is None or r['id'] in selected]
    headers=['文件名','证书类型','审核状态','期限提示']+[label for _,label in FIELD_DEFS]+['缺失必填字段','待核对说明','OCR耗时秒','审核人','审核时间','原始识别记录ID']
    records=[]
    headers+=['上次复审日期','下次应复审日期','复审证明状态','复审证明依据','复审风险']
    headers+=['人员编号','档案姓名','所属项目','岗位','负责安全员','证书版本状态']
    statuslabels={'missing':'待补录','review':'待核对','ready':'待审核','approved':'已审核'}
    for r in rows:
        vals=[r['filename'],r['type_label'],statuslabels[r['status']] if r['job_status']=='done' else r['job_status'],r['risk']['label']]
        for key,_ in FIELD_DEFS:
            f=r['fields'][key]
            vals.append('不适用' if not f['applicable'] else f['value'] or ('【未识别】' if f['required'] else ''))
        issues=[f'{LABELS[k]}：'+ '；'.join(r['fields'][k]['issues']) for k in r['needs_review']]+r['notes']
        vals.extend(['、'.join(LABELS[k] for k in r['missing']),'；'.join(issues),r['ocr_s'],r['reviewer'],r['reviewed_at'] or '',r['id']])
        info=r['review_info'];applicable=r['type']=='electrical'
        vals.extend([info['last_review_date'],info['next_due_date'],{'unknown':'待核实','pending':'待补充','checked':'人工已核对'}[info['proof_status']],info['proof_note'],r['review_risk']['label']] if applicable else ['不适用']*5)
        person=people.get(r.get('person_id'),{})
        vals.extend([person.get(k,'') for k in ['id','name','project','role','officer']]+['历史证书' if r['id'] in historical else '当前证书'])
        records.append(vals)
    if body.get('format')=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill,Alignment
        from openpyxl.utils import get_column_letter
        wb=Workbook(); ws=wb.active; ws.title='证书提取结果'
        ws.append(headers)
        for row in records: ws.append([spreadsheet_safe(v) for v in row])
        for c in ws[1]: c.fill=PatternFill('solid',fgColor='155E68'); c.font=Font(color='FFFFFF',bold=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment=Alignment(vertical='center',wrap_text=True)
                if cell.value=='【未识别】': cell.fill=PatternFill('solid',fgColor='FDE8E7'); cell.font=Font(color='B33131')
                elif cell.value=='不适用': cell.font=Font(color='929A9E')
        for col in range(1,len(headers)+1): ws.column_dimensions[get_column_letter(col)].width=25 if col!=1 else 42
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        output=io.BytesIO(); wb.save(output)
        return Response(output.getvalue(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="certificate-results.xlsx"'})
    output=io.StringIO(newline=''); writer=csv.writer(output); writer.writerow(headers)
    writer.writerows([[spreadsheet_safe(v) for v in row] for row in records])
    return Response(output.getvalue().encode('utf-8-sig'),media_type='text/csv; charset=utf-8',headers={'Content-Disposition':'attachment; filename="certificate-results.csv"'})

@app.get('/')
def index(): return FileResponse(ROOT/'static'/'index.html')
app.mount('/static',StaticFiles(directory=ROOT/'static'),name='static')
operations.register(sys.modules[__name__])
invoice_service.register(sys.modules[__name__])

if __name__=='__main__':
    p=argparse.ArgumentParser(); p.add_argument('--port',type=int,default=8765); args=p.parse_args()
    import uvicorn
    uvicorn.run(app,host='127.0.0.1',port=args.port,log_level='info')
